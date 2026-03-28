# ./venv/Scripts/python.exe automated_ale_writer.py

import os
import glob
import pandas as pd
import re
from bs4 import BeautifulSoup
from pptx import Presentation
from langchain_community.tools import DuckDuckGoSearchRun
from langchain_core.output_parsers import StrOutputParser
from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from langchain_ollama import ChatOllama
from langchain_classic.chains import RetrievalQA
from langchain_core.prompts import PromptTemplate
from langchain_core.documents import Document

# --- 1. CONFIGURATION ---
DATA_FOLDER = "./RAG - ALE ARCHITECTURE BOARD EXAM"
OUTPUT_FOLDER = "./ALE_Reviewer_Library"
DB_DIR = "./chroma_db"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# THE TITAN DICTIONARY (Comprehensive ALE Gaps)
# TARGET_CONTENT = {
#     "Accessibility Standards (BP 344 - RA 9442)": [
#         "BP 344: Minimum Requirements for Curb Cut-outs and Ramps (1:12 Slope)",
#         "BP 344: Standards for Accessible Washrooms, Toilets, and Grab Bars",
#         "BP 344: Tactile Blocks, Braille Signage, and Visual Warning Signals",
#         "RA 9442: Privileges and Incentives for PWDs in Building Design",
#         "Dimensions for Wheelchair Maneuvering and PWD Parking Slots"
#     ],
#     "National Building Code Application (PD 1096)": [
#         "PD 1096 Rule VII: Classification and General Requirements of all Buildings",
#         "PD 1096 Rule VIII: Light and Ventilation (Window Sizes and Court Dimensions)",
#         "PD 1096 Rule XII: General Requirements for Egress and Occupant Loads",
#         "Calculating GFA, TGFA, PSO, and Building Height Limits (BHL)",
#         "Building Permit Process, Fines, and the Role of the Building Official"
#     ],
#     "Housing & Human Settlements (BP 220 - PD 957)": [
#         "BP 220: Technical Standards for Socialized and Economic Housing",
#         "PD 957: Subdivision and Condominium Buyer's Protective Decree",
#         "Comparison of Road Widths and Open Space: BP 220 vs PD 957",
#         "Housing Policies and Relocation Strategies for Informal Settlements"
#     ],
#     "Anthropometrics & Human Dimensions": [
#         "Anthropometrics: Standard Human Body Dimensions (Standing, Sitting, Reaching)",
#         "Residential Clearances: Kitchen Work Triangles and Bedroom Layouts",
#         "Commercial Clearances: Office Workstations and Retail Aisle Widths",
#         "Human Dimensions for Institutional Design (Schools and Hospitals)"
#     ],
#     "Planning: Environmental Factors & Site Analysis": [
#         "Microclimate Analysis: Wind Paths (Amihan/Habagat) and Solar Orientation",
#         "Topography and Slope Analysis: Cut-and-Fill and Buildability Limits",
#         "Geological Hazards: Identifying Fault Lines, Liquefaction, and Flood Zones",
#         "Hydrology in Site Planning: Stormwater Runoff and Drainage Systems"
#     ],
#     "History: Philippine Architecture & Identity": [
#         "Pre-Colonial Vernacular: The Evolution of the Bahay Kubo",
#         "Spanish Colonial Era: The Bahay na Bato and Earthquake Baroque",
#         "American Colonial Era: Daniel Burnham’s Manila and the Pensionado Architects",
#         "Masters of the Golden Age: Juan Nakpil, Pablo Antonio, and Leandro Locsin",
#         "The Search for Modern Filipino Identity: Mañosa and the New Vernacular"
#     ],
#     "ALE Mastery & Study Strategies": [
#         "Architectural Design & Site Planning: Mastery and Study Strategy",
#         "Building Utilities & MEPF: Integrated Review and Connections",
#         "Building Materials & Construction: Common Board Exam Materials",
#         "Professional Practice & Jurisprudence: The 15-Year Liability Framework",
#         "Structural Conceptualization: Load Paths and Seismic Zone 4 Rules"
#     ]
# }

TARGET_CONTENT = {
    "Accessibility Standards (BP 344 - RA 9442)": [
        "BP 344 (2024 IRR): Ramp Slope Tiers (1:20, 1:15, 1:12) and Maximum Run Lengths",
        "BP 344 (2024 IRR): Accessible Washrooms – Clearance, Grab Bars, and Fixture Heights",
        "BP 344 (2024 IRR): Tactile Warning Surfaces, Braille Signage, and Visual Alarms",
        "RA 9442: Mandatory PWD Privileges, Incentives, and Design Integration",
        "Dimensions for Wheelchair Maneuvering, Parking Slots (4.00m min.), and Passing Spaces"
    ],
    "National Building Code Application (PD 1096)": [
        "PD 1096 Rule VII: Classification and General Requirements of all Buildings",
        "PD 1096 Rule VIII: Light and Ventilation (Window Sizes and Court Dimensions)",
        "PD 1096 Rule XII: General Requirements for Egress and Occupant Loads",
        "Calculating GFA, TGFA, PSO, and Building Height Limits (BHL)",
        "Building Permit Process, Fines, and the Role of the Building Official"
    ],
    "Housing & Human Settlements (BP 220 - PD 957)": [
        "BP 220: Technical Standards for Socialized and Economic Housing",
        "PD 957: Subdivision and Condominium Buyer's Protective Decree",
        "Comparison of Road Widths and Open Space: BP 220 vs PD 957",
        "Housing Policies and Relocation Strategies for Informal Settlements"
    ],
    "Anthropometrics & Human Dimensions": [
        "Anthropometrics: Standard Human Body Dimensions (Standing, Sitting, Reaching)",
        "Residential Clearances: Kitchen Work Triangles and Bedroom Layouts",
        "Commercial Clearances: Office Workstations and Retail Aisle Widths",
        "Human Dimensions for Institutional Design (Schools and Hospitals)"
    ],
    "Planning: Environmental Factors & Site Analysis": [
        "Microclimate Analysis: Wind Paths (Amihan/Habagat) and Solar Orientation",
        "Topography and Slope Analysis: Cut-and-Fill and Buildability Limits",
        "Geological Hazards: Identifying Fault Lines, Liquefaction, and Flood Zones",
        "Hydrology in Site Planning: Stormwater Runoff and Drainage Systems"
    ],
    "History: Philippine Architecture & Identity": [
        "Pre-Colonial Vernacular: The Evolution of the Bahay Kubo",
        "Spanish Colonial Era: The Bahay na Bato and Earthquake Baroque",
        "American Colonial Era: Daniel Burnham’s Manila and the Pensionado Architects",
        "Masters of the Golden Age: Juan Nakpil, Pablo Antonio, and Leandro Locsin",
        "The Search for Modern Filipino Identity: Mañosa and the New Vernacular"
    ],
    "ALE Mastery & Study Strategies": [
        "Architectural Design & Site Planning: Mastery and Study Strategy",
        "Building Utilities & MEPF: Integrated Review and Connections",
        "Building Materials & Construction: Common Board Exam Materials",
        "Professional Practice & Jurisprudence: The 15-Year Liability Framework",
        "Structural Conceptualization: Load Paths and Seismic Zone 4 Rules"
    ]
}

# --- 2. GPU EMBEDDINGS ---
print("Waking up the RTX 4070...")
embeddings = HuggingFaceEmbeddings(
    model_name="all-MiniLM-L6-v2",
    model_kwargs={'device': 'cuda'}
)

# --- 3. SMART LOADING (SQL Paginated Fix) ---
existing_sources = set()
if os.path.exists(DB_DIR):
    print("Database found! Safely fetching existing index (Paginated)...")
    vectorstore = Chroma(persist_directory=DB_DIR,
                         embedding_function=embeddings)
    try:
        count = vectorstore._collection.count()
        offset, limit = 0, 5000
        while offset < count:
            batch_data = vectorstore._collection.get(
                include=['metadatas'], limit=limit, offset=offset)
            for meta in batch_data['metadatas']:
                if meta and 'source' in meta:
                    existing_sources.add(meta['source'])
            offset += limit
        print(
            f" [+] Verified {len(existing_sources)} existing source files in the brain.")
    except Exception as e:
        print(f" [!] Metadata fetch issue: {e}")
else:
    vectorstore = None

# --- 4. DOCUMENT LOADING & CLEANING ---
print("Scanning for documents...")
new_docs = []
all_files = glob.glob(os.path.join(DATA_FOLDER, "**", "*.*"), recursive=True)
for file_path in all_files:
    abs_path = os.path.abspath(file_path)
    if abs_path in existing_sources:
        continue
    ext = abs_path.lower().split('.')[-1]
    try:
        if ext == 'pdf':
            loader = PyPDFLoader(abs_path)
            new_docs.extend([d for d in loader.load() if d.page_content and len(
                d.page_content.strip()) > 5])
        elif ext == 'docx':
            loader = Docx2txtLoader(abs_path)
            new_docs.extend(loader.load())
        elif ext == 'pptx':
            prs = Presentation(abs_path)
            text = " ".join(
                [shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")])
            if len(text.strip()) > 10:
                new_docs.append(Document(page_content=text,
                                metadata={"source": abs_path}))
        print(f" [+] New Data Found: {os.path.basename(abs_path)}")
    except Exception:
        continue

# --- 5. UNSTOPPABLE BATCHER ---
if new_docs:
    print(f"Slicing {len(new_docs)} new sections...")
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000, chunk_overlap=100)
    valid_chunks = [c for c in text_splitter.split_documents(new_docs) if isinstance(
        c.page_content, str) and len(c.page_content.strip()) > 10]
    batch_size = 500
    for i in range(0, len(valid_chunks), batch_size):
        batch = valid_chunks[i: i + batch_size]
        try:
            if not vectorstore:
                vectorstore = Chroma.from_documents(
                    documents=batch, embedding=embeddings, persist_directory=DB_DIR)
            else:
                vectorstore.add_documents(batch)
            print(f"  [#] Batch {i//batch_size + 1} indexed...")
        except Exception:
            for chunk in batch:
                try:
                    vectorstore.add_documents([chunk])
                except:
                    continue
else:
    print("Database up to date! Jumping to generation.")

# --- 6. "REVIEWER PRO" PROMPT & LLM ---
llm = ChatOllama(model="qwen3.5:9b", temperature=0.2)
web_search = DuckDuckGoSearchRun()  # Initialize the Internet connection
professional_prompt = """
You are a fused persona: a Senior Architectural Reviewer and a World-Class Educator for the Philippine ALE. Your single goal is to produce a "Masterclass Study Module" that ensures both technical mastery and deep conceptual retention, preparing the user not just to pass, but to excel. You are a master of storytelling, using analogies to simplify complex topics, and a master of the exam, knowing exactly how the Board of Architecture will test this subject.

CRITICAL DIRECTIVE: You have been provided with context containing both a "LOCAL 5GB DATABASE" and a "LIVE INTERNET SEARCH". You must actively cross-reference these. If the live internet data shows a recent amendment, a newly verified standard, or an updated Implementing Rules and Regulations (IRR), you MUST prioritize the updated internet data over older local data. 

Target Topic: {question}

Output Format: Generate the response using the following 9-part structure. Use Markdown for formatting, including headings, tables, and bullet points.

### 1. The "Big Picture" & ALE Focus
- **The Hook:** Begin with 2-3 sentences that frame the topic as a solution to a real-world problem. Use a powerful analogy (e.g., "Think of the stair as a building's circulatory system...").
- **ALE Importance:** Clearly state why the Board of Architecture cares about this topic. What specific skill or competency are they testing (e.g., "life safety," "spatial efficiency," "legal compliance")?

### 2. Narrative Logic & First Principles (The "Why")
- Explain the logic, history, or intent behind the rules. This section is about deep understanding, not rote memorization.
    - **If a Law (e.g., PD 1096):** Explain the "Intent of the Law" (e.g., "Fire safety is paramount, so the code prioritizes compartmentation and egress...").
    - **If History:** Explain how climate, available technology, or socio-political power structures shaped the architectural form.
    - **If Construction:** Explain how the material behaves under stress (tension/compression) and how the construction method responds to that behavior.

### 3. Technical Specifications & Spatial Data (The "What")
- Provide a high-density **Markdown Table** of all "Must-Memorize" dimensions, percentages, ratios, or dates. Structure the table for easy recall.
- Use **LaTeX** for all formulas and mathematical expressions: e.g., $Slope = \\frac{{Rise}}{{Run}} \\times 100$ or $GFA = Area_{{Floors}} + Area_{{Mezzanines}}$.
- Ensure all measurements are in Philippine Standard units (mm, m, sqm).
- **Latest Legal Updates:** Explicitly flag any new amendments or recent IRR updates found in the live internet context (e.g., noting if a rule is from the newer 2024 IRR vs older versions). 

### 4. Computation & Application Patterns (The "How")
- Provide 1-2 practical, exam-style problems based on the context.
- **Difficulty Rating:** Include a difficulty rating from 0.1 (basic) to 1.0 (expert-level).
- **Step-by-Step Solution:** Provide a clear, step-by-step solution. Use LaTeX to show the calculation process so the student can follow the math.

### 5. The Examiner's Mindset (Trap Analysis)
- Identify the specific ways this topic is "weaponized" in the board exam. This is your "inside information."
- **Pattern 1: The 'Trick' Question:** Describe how they phrase the question to confuse the candidate (e.g., mixing units, providing irrelevant data, or asking for the *exception* to a rule).
- **Pattern 2: The Red Herring:** List 3 common "Distractor" answers that look correct but are intentionally wrong, and explain *why* they are wrong.

### 6. How It's Tested on the ALE (Question Patterns)
- List 3-4 specific question patterns found in board exams. Use a clear format:
    - **Pattern: [Question Type, e.g., "Scenario-Based"]**
    - **Example:** "In a 5-story office building, the required fire escape stair width is calculated based on..."
    - **Focus:** "Tests your knowledge of occupancy load vs. egress capacity."

### 7. Active Recall Challenge (Flashcards)
- Provide 3 "Flashcard" questions that require deep thinking, not just memorization.
    - Include one **"Complex Application"** question where the student must synthesize this topic with another (e.g., "How does the minimum lot area from BP 220 affect the required fire lane widths from RA 9514?").

### 8. The Web of Knowledge (Interconnectedness)
- Explain how this topic connects to other ALE subjects. This helps build a holistic mental map.
    - **Example:** "Understanding the history of the *Bahay Kubo* (History) is foundational to grasping the passive cooling and ventilation requirements in Rule VIII of the NBC (Building Technology)."

### 9. Key Terms & Sources
- **Definitions:** 5-8 essential technical terms with concise, board-exam-relevant definitions.
- **Verification:** Cite specific Philippine Laws (e.g., PD 1096, BP 344, RA 9266, BP 220, etc.) and any relevant source documents provided in the context.
- **Live Updates:** If the internet context provided newer information that overrode the local database, explicitly state: *"Update Note: Verified against latest internet search..."*

---
CONTEXT: {context}
QUERY: {question}
MASTERCLASS MODULE:
"""
PROMPT = PromptTemplate(template=professional_prompt,
                        input_variables=["context", "question"])

# Use modern LCEL syntax instead of RetrievalQA
hybrid_chain = PROMPT | llm | StrOutputParser()

# --- 7. HYBRID GENERATION & INTEGRATED INDEXING ---
print("\nWRITING CONNECTED MASTERCLASS ARTICLES...")
article_stats = []

for subtopic, articles in TARGET_CONTENT.items():
    clean_subfolder = subtopic.replace(
        ":", " -").replace("/", "-").replace("?", "").replace('"', "")
    subfolder_path = os.path.join(OUTPUT_FOLDER, clean_subfolder)
    os.makedirs(subfolder_path, exist_ok=True)

    for title in articles:
        safe_name = title.replace(
            ":", " -").replace("/", "-").replace("?", "").replace('"', "") + ".md"
        file_path = os.path.join(subfolder_path, safe_name)

        if not os.path.exists(file_path):
            print(f"\n[Search] Gathering data for: {title}...")

            # 1. Fetch from Local 5GB Brain
            local_context = ""
            if vectorstore:
                local_docs = vectorstore.similarity_search(title, k=10)
                local_context = "\n".join(
                    [doc.page_content for doc in local_docs])

            # 2. Fetch from the Live Internet
            print("  [🌐] Pinging DuckDuckGo for live context...")
            try:
                # Add "Philippine Architecture" to focus the search engine
                internet_context = web_search.invoke(
                    title + " Philippine Architecture")
            except Exception as e:
                internet_context = "Internet search failed or timed out."

            # 3. Combine them into the Ultimate Context
            combined_context = f"--- LOCAL 5GB DATABASE ---\n{local_context}\n\n--- LIVE INTERNET SEARCH ---\n{internet_context}"

            print(f"  [✍️] Writing module...")
            try:
                content = hybrid_chain.invoke(
                    {"context": combined_context, "question": title})
                with open(file_path, "w", encoding="utf-8") as f:
                    f.write(content)
            except Exception as e:
                print(f"  [!] Error: {e}")
                content = ""
        else:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()

        # Audit stats for the Master Index
        word_count = len(content.split())
        status = "✅ Complete" if word_count > 300 else "⚠️ Thin (Needs More Data)"
        article_stats.append({"Subtopic": clean_subfolder, "Title": title, "Words": word_count,
                              "Status": status, "Path": f"./{clean_subfolder}/{safe_name}"})

# --- 8. GENERATE MASTER INDEX (INTEGRATED FEATURE) ---
index_path = os.path.join(OUTPUT_FOLDER, "MASTER_INDEX.md")
print(f"\nGenerating Master Index at {index_path}...")

with open(index_path, "w", encoding="utf-8") as f:
    f.write("# 🏛️ ALE ARCHITECTURAL REVIEWER MASTER INDEX\n")
    f.write(f"*Total Articles Generated: {len(article_stats)}*\n\n")
    f.write("| Subtopic | Article Title | Word Count | Status | Link |\n")
    f.write("| :--- | :--- | :--- | :--- | :--- |\n")
    for stat in article_stats:
        f.write(
            f"| {stat['Subtopic']} | {stat['Title']} | {stat['Words']} | {stat['Status']} | [View Article]({stat['Path']}) |\n")

print("\nMISSION COMPLETE. Review your library and MASTER_INDEX.md in ./ALE_Reviewer_Library")
